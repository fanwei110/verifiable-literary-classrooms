"""将 DOCX 中所有表格设置为页面正文宽度，并恢复三线表。

本脚本用于投稿前排版后处理。Pandoc 可以把 Markdown 管道表转换为
Word 表格，但表格宽度通常由内容自动决定；这里统一写入 Word 的
100% 首选宽度，使表格在 PDF 导出时铺满正文版心。同时显式写入
三线表边框，避免 Word 默认表格网格覆盖投稿稿件的学术表格样式。
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_table_width_100_percent(table) -> None:
    """把单个 Word 表格设为 100% 正文宽度。

    Parameters
    ----------
    table : docx.table.Table
        需要调整宽度的 Word 表格对象。

    Notes
    -----
    使用百分比宽度而非固定厘米值，是为了让同一文件在不同模板、
    页边距或后续 IEEE Access 模板迁移时仍能跟随正文版心变化。
    """

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "pct")
    tbl_w.set(qn("w:w"), "5000")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "autofit")


def _get_or_add(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def _set_border(parent, side: str, val: str, size: str = "8") -> None:
    border = _get_or_add(parent, f"w:{side}")
    border.set(qn("w:val"), val)
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), "0")
    border.set(qn("w:color"), "000000")


def _set_cell_borders(cell, borders: dict[str, str]) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = _get_or_add(tc_pr, "w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        _set_border(tc_borders, side, borders.get(side, "nil"))


def set_three_line_table_borders(table) -> None:
    """把单个 Word 表格设置为三线表。

    Parameters
    ----------
    table : docx.table.Table
        需要调整边框的 Word 表格对象。

    Notes
    -----
    三线表只保留表顶线、表头分隔线和表底线。先清掉单元格四周边框，
    再对首行和末行补线，是为了压过 Word/Pandoc 可能继承的网格样式。
    """

    if not table.rows:
        return

    tbl_pr = table._tbl.tblPr
    tbl_borders = _get_or_add(tbl_pr, "w:tblBorders")
    for side in ("left", "right", "insideH", "insideV"):
        _set_border(tbl_borders, side, "nil")
    _set_border(tbl_borders, "top", "single", "8")
    _set_border(tbl_borders, "bottom", "single", "8")

    for row in table.rows:
        for cell in row.cells:
            _set_cell_borders(cell, {})

    for cell in table.rows[0].cells:
        _set_cell_borders(cell, {"top": "single", "bottom": "single"})

    for cell in table.rows[-1].cells:
        _set_cell_borders(cell, {"bottom": "single"})


def main() -> int:
    """读取输入 DOCX，设置所有表格宽度，并写出结果文件。"""

    if len(sys.argv) != 3:
        print("Usage: python set_docx_table_full_width.py INPUT.docx OUTPUT.docx")
        return 2

    src = Path(sys.argv[1])
    dst = Path(sys.argv[2])
    doc = Document(src)

    for table in doc.tables:
        set_table_width_100_percent(table)
        set_three_line_table_borders(table)

    doc.save(dst)
    print(f"updated {len(doc.tables)} tables -> {dst}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
